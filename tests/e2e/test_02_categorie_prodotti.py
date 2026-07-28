import pytest
import os
from openpyxl import load_workbook
from playwright.sync_api import Page, expect
from utils.logger import log_scoperta

# Percorso all'excel fornito dall'utente
import os
EXCEL_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../audit/MaterialePerTest/Prodotti_Tracciato.xlsx"))

def get_categorie_da_excel():
    """Estrae l'elenco delle categorie uniche dal file Excel"""
    if not os.path.exists(EXCEL_PATH):
        return []
        
    try:
        wb = load_workbook(EXCEL_PATH, data_only=True)
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]
        if 'categoria' in headers:
            cat_idx = headers.index('categoria')
            categories = set()
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row[cat_idx]:
                    categories.add(row[cat_idx])
            return list(categories)
    except Exception:
        pass
    return []

def test_creazione_categorie_da_excel(logged_page: Page):
    """
    Testa la creazione delle categorie prodotti leggendole dal file Excel fornito.
    """
    categorie_excel = get_categorie_da_excel()
    if not categorie_excel:
        pytest.skip("File Excel non trovato o nessuna categoria presente.")
        
    page = logged_page
    
    # Navigazione alla pagina Categorie Prodotti
    page.goto("products/categories")
    
    # Verifica il caricamento della pagina
    expect(page.locator("text=Elenco Categorie")).to_be_visible()
    
    for nome_cat in categorie_excel:
        # Clicca "Nuova Categoria"
        page.click("button:has-text('Nuova Categoria')")
        
        # Attendi che il modale si apra
        expect(page.locator("#modal-categoria")).to_be_visible()
        
        # Compila il nome della categoria
        page.fill("#cat-nome", nome_cat)
        
        # Test della categoria genitore con QkM2O (non deterministico, ~30% probabilità)
        import random
        import time
        if random.random() < 0.3 and categorie_excel:
            parent_cat = random.choice(categorie_excel)
            if parent_cat != nome_cat: # Evita di mettersi come genitore di se stesso
                parent_input = page.locator("#cat-parent-label")
                parent_input.click()
                parent_input.fill("")
                time.sleep(0.5)
                parent_input.type(parent_cat, delay=100)
                
                dropdown_opt = page.locator(f"#cat-parent-dropdown .qk-m2o-option:has-text('{parent_cat}')").first
                try:
                    dropdown_opt.wait_for(state="visible", timeout=5000)
                    dropdown_opt.click()
                except Exception:
                    pass # Se non la trova (magari non è ancora stata creata), proseguiamo
        
        # Salva
        page.click("#modal-categoria button[type='submit']")
        
        try:
            # Aspettiamo la scomparsa del modale o un segnale di refresh/successo
            # In questo caso potremmo avere un redirect o un reload
            page.wait_for_load_state('networkidle')
            expect(page.locator(f"text='{nome_cat}'").first).to_be_visible(timeout=5000)
            
            log_scoperta(
                modulo="Categorie Prodotti", 
                test_name="Creazione Categoria",
                descrizione=f"Categoria '{nome_cat}' creata con successo (estratta da Excel).",
                dati_usati={"nome_categoria": nome_cat}
            )
        except Exception as e:
            # Se la categoria esiste gi, o c' un errore
            page.screenshot(path=f"test-results/categoria_{nome_cat.replace(' ', '_')}_fail.png")
            log_scoperta(
                modulo="Categorie Prodotti", 
                test_name="Creazione Categoria - FALLITO",
                descrizione=f"Impossibile creare categoria '{nome_cat}'. {str(e)}",
                dati_usati={"nome_categoria": nome_cat},
                screenshot_path=f"test-results/categoria_{nome_cat.replace(' ', '_')}_fail.png"
            )
            # Continuiamo con le altre categorie anche se una fallisce
            # (es. potrebbe essere gi esistente)
            # Ricarichiamo la pagina per resettare lo stato
            page.goto("products/categories")

